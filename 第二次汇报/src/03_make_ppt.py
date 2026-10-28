"""
生成第二次汇报 PPT:
  输出: ../Jura_MOGP_第二次汇报.pptx

严格按照 ../PPT规划.md、../方案.md、../结果.md 编写:
  正文 9 页 + 致谢 1 页 + 备用页 4 页 (共 14 页)
  延续第一次汇报 (Jura_MOGP_组会汇报.pptx) 的 16:9 宽屏排版风格与配色体系。
"""

from pathlib import Path
from PIL import Image
import pandas as pd
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).resolve().parent.parent          # E:\Research\第二次汇报
FIG = HERE / "figures"
RES = HERE / "results"
OUT_PPTX = HERE / "Jura_MOGP_第二次汇报.pptx"

# ---------- 配色体系 (延续第一次汇报 + 适配本次图表) ----------
C_DARK   = RGBColor(0x1F, 0x29, 0x37)   # 深灰蓝 (主标题、表头、页码徽章)
C_CARD_D = RGBColor(0x27, 0x33, 0x4A)   # 封面深色卡片
C_TEXT   = RGBColor(0x1F, 0x29, 0x37)   # 正文深色
C_MUTED  = RGBColor(0x6B, 0x72, 0x80)   # 辅助灰字
C_LIGHT  = RGBColor(0x9C, 0xA3, 0xAF)   # 页脚浅灰
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

C_BLUE   = RGBColor(0x2A, 0x78, 0xD6)   # 主强调蓝 (与图1一致)
C_STEEL  = RGBColor(0x4C, 0x72, 0xB0)   # 稳重蓝
C_RED    = RGBColor(0xC0, 0x4A, 0x3A)   # 警示/重点红
C_GREEN  = RGBColor(0x1E, 0x7E, 0x5A)   # 达标绿

BG_GRAY  = RGBColor(0xF3, 0xF4, 0xF6)   # 浅灰底色框
BG_BLUE  = RGBColor(0xE8, 0xF0, 0xFE)   # 浅蓝提示框
BG_RED   = RGBColor(0xF9, 0xEB, 0xEA)   # 浅红重点框
BG_GREEN = RGBColor(0xE9, 0xF7, 0xEF)   # 浅绿提示框
BORDER_G = RGBColor(0xD1, 0xD5, 0xDB)

FONT_HEAD = "Microsoft YaHei"
FONT_BODY = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

TOTAL_MAIN = 9


def set_bg(slide, color=C_WHITE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_box(slide, left, top, width, height, fill_color=None, line_color=None, line_pt=1.0, rounded=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(
        shape_type, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    if rounded and sh.adjustments:
        sh.adjustments[0] = 0.04
    if fill_color is not None:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill_color
    else:
        sh.fill.background()
    if line_color is not None:
        sh.line.color.rgb = line_color
        sh.line.width = Pt(line_pt)
    else:
        sh.line.fill.background()
    return sh


def add_text(slide, left, top, width, height, text, size=14, color=C_TEXT,
             bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             font_name=FONT_BODY, line_spacing=1.25):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)

    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if i > 0:
            p.space_before = Pt(4)
        run = p.add_run()
        run.text = line
        run.font.name = font_name
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


def add_rich_paragraphs(slide, left, top, width, height, items,
                        valign=MSO_ANCHOR.TOP, line_spacing=1.25, space_after=6):
    """
    items: list of paragraphs, each paragraph is a list of (text, size, color, bold) tuples.
    """
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)

    for i, para_runs in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        for txt, sz, col, bld in para_runs:
            r = p.add_run()
            r.text = txt
            r.font.name = FONT_BODY
            r.font.size = Pt(sz)
            r.font.color.rgb = col
            r.font.bold = bld
    return tb


def add_header(slide, badge_str, title_str, subtitle_str=None, is_backup=False):
    badge_bg = C_STEEL if is_backup else C_DARK
    add_box(slide, 0.50, 0.42, 0.56, 0.50, fill_color=badge_bg)
    add_text(slide, 0.50, 0.42, 0.56, 0.50, badge_str, size=14, color=C_WHITE,
             bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 1.20, 0.36, 11.60, 0.58, title_str, size=24, color=C_DARK,
             bold=True, font_name=FONT_HEAD)
    # 细分割线
    add_box(slide, 0.50, 1.05, 12.33, 0.02, fill_color= RGBColor(0xE5, 0xE7, 0xEB))
    if subtitle_str:
        add_text(slide, 0.50, 1.14, 12.33, 0.40, subtitle_str, size=13.5, color=C_MUTED)


def add_footer(slide, page_label, dark=False):
    col = C_LIGHT if dark else C_MUTED
    add_text(slide, 0.50, 7.12, 6.50, 0.28,
             "Jura 土壤镉浓度预测 · 第二次汇报（公平比较 MOGP 与线性回归）",
             size=9.5, color=col)
    add_text(slide, 11.30, 7.12, 1.53, 0.28, page_label,
             size=9.5, color=col, align=PP_ALIGN.RIGHT)


def add_callout(slide, left, top, width, height, title, body,
                bg_color=BG_BLUE, accent_color=C_BLUE, title_size=13, body_size=12.5):
    add_box(slide, left, top, width, height, fill_color=bg_color)
    add_box(slide, left, top, 0.08, height, fill_color=accent_color)
    add_rich_paragraphs(
        slide, left + 0.22, top + 0.10, width - 0.34, height - 0.18,
        [
            [(title, title_size, accent_color, True)],
            [(body, body_size, C_TEXT, False)],
        ],
        line_spacing=1.22, space_after=3
    )


def add_picture_fit(slide, img_path, box_left, box_top, box_w, box_h):
    """按原始宽高比居中放入指定矩形区域, 绝不拉伸变形。"""
    with Image.open(img_path) as im:
        w_px, h_px = im.size
    aspect = w_px / h_px
    box_aspect = box_w / box_h
    if aspect >= box_aspect:
        w = box_w
        h = box_w / aspect
    else:
        h = box_h
        w = box_h * aspect
    left = box_left + (box_w - w) / 2
    top = box_top + (box_h - h) / 2
    return slide.shapes.add_picture(str(img_path), Inches(left), Inches(top),
                                    width=Inches(w), height=Inches(h))


def style_table(table, col_widths, rows_data, alignments=None,
                header_bg=C_DARK, highlight_rows=None, font_size=12.5, header_size=12.5):
    """
    rows_data: list of rows (row 0 is header). Each cell can be str or (str, color, bold).
    highlight_rows: dict {row_idx: bg_color}
    """
    highlight_rows = highlight_rows or {}
    for c_i, w in enumerate(col_widths):
        table.columns[c_i].width = Inches(w)

    for r_i, row in enumerate(rows_data):
        is_header = (r_i == 0)
        default_bg = header_bg if is_header else (
            highlight_rows.get(r_i, BG_GRAY if r_i % 2 == 1 else C_WHITE)
        )
        for c_i, val in enumerate(row):
            cell = table.cell(r_i, c_i)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.10)
            cell.margin_right = Inches(0.10)
            cell.margin_top = Inches(0.06)
            cell.margin_bottom = Inches(0.06)
            cell.fill.solid()
            cell.fill.fore_color.rgb = default_bg

            if isinstance(val, tuple):
                txt, col, bld = val
            else:
                txt = str(val)
                col = C_WHITE if is_header else C_TEXT
                bld = True if (is_header or c_i == 0) else False

            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = (alignments[c_i] if alignments else
                           (PP_ALIGN.LEFT if c_i == 0 else PP_ALIGN.CENTER))
            p.line_spacing = 1.15
            run = p.add_run()
            run.text = txt
            run.font.name = FONT_BODY
            run.font.size = Pt(header_size if is_header else font_size)
            run.font.color.rgb = col
            run.font.bold = bld


# ==============================================================================
# 第 1 页：标题页
# ==============================================================================
s1 = prs.slides.add_slide(BLANK)
set_bg(s1, C_DARK)

add_text(s1, 0.90, 1.55, 11.50, 0.40,
         "组会汇报（第二次） · 瑞士 JURA 土壤重金属数据集",
         size=13.5, color=RGBColor(0x9D, 0xB2, 0xD9), bold=True)

add_text(s1, 0.90, 2.05, 11.50, 1.10,
         "公平比较 MOGP 与线性回归",
         size=40, color=C_WHITE, bold=True, font_name=FONT_HEAD)

add_text(s1, 0.90, 3.20, 11.20, 0.85,
         "上次汇报被指出“偏袒 MOGP、线性回归对比草率”。\n"
         "这次只做一件事：在赛前定死规则、掌握同样信息的前提下，检验“参考邻居”到底有没有用。",
         size=16, color=RGBColor(0xCB, 0xD5, 0xE1), line_spacing=1.35)

add_box(s1, 0.90, 4.30, 3.50, 0.04, fill_color=C_STEEL)
add_text(s1, 0.90, 4.50, 9.00, 0.40,
         "汇报人：蔡昊璋（Aza）      2026 年 10 月",
         size=13, color=C_WHITE)

# 底部三张核心结论速览卡片
cards_s1 = [
    ("6 条公平规则", "同一套 10 次划分 + 赛前定死 p < 0.05 判定标准"),
    ("+6.4% (p=0.0001)", "主比较：MOGP 胜过同信息的回归 B（10/10 全胜）"),
    ("p = 0.33（打平）", "附加比较：回归 C 加上岩石/土地用途即可追平 MOGP"),
]
for idx, (big_t, sub_t) in enumerate(cards_s1):
    x0 = 0.90 + idx * 3.90
    add_box(s1, x0, 5.25, 3.65, 1.35, fill_color=C_CARD_D, rounded=True)
    add_text(s1, x0 + 0.18, 5.38, 3.30, 0.55, big_t,
             size=21, color=RGBColor(0x7F, 0xA8, 0xE0), bold=True)
    add_text(s1, x0 + 0.18, 5.98, 3.30, 0.50, sub_t,
             size=11.5, color=RGBColor(0xCB, 0xD5, 0xE1))

add_footer(s1, f"1 / {TOTAL_MAIN}", dark=True)


# ==============================================================================
# 第 2 页：上次的问题和这次的改进
# ==============================================================================
s2 = prs.slides.add_slide(BLANK)
set_bg(s2)
add_header(s2, "02", "上次汇报的问题与这次的改进",
           "上次被指出“完全在偏袒 MOGP，线性回归对比做得草率”——逐条复盘并彻底改正")

tb2 = s2.shapes.add_table(5, 3, Inches(0.50), Inches(1.65), Inches(12.33), Inches(3.85)).table
style_table(
    tb2,
    col_widths=[2.53, 4.70, 5.10],
    rows_data=[
        ["上次的不公平之处", "第一次汇报的做法（偏袒 MOGP）", "第二次汇报的改进（公平对决）"],
        [
            "1. 考的次数不一样",
            "MOGP 考了 10 次随机划分，线性回归只考了 1 次",
            ("四个模型考完全相同的 10 次随机划分（同一张考卷）", C_BLUE, True)
        ],
        [
            "2. 没有做显著性检验",
            "只报了“GP 比回归多降 4% 误差”，未检验是不是运气",
            ("每次划分算配对差值，做配对 t 检验 + Wilcoxon 检验", C_BLUE, True)
        ],
        [
            "3. 头条数字比错了对象",
            "“+26%”比的是完全不看 Ni/Zn 的单输出 GP，夸大了优势",
            ("主比较对手换成同样知道 Ni、Zn 和位置的回归 B", C_BLUE, True)
        ],
        [
            "4. 优点只说不测（说错了）",
            "声称“线性回归给不出不确定度”（其实课本里就有公式）",
            ("回归和 MOGP 都交出 95% 预测区间，实测对比覆盖率", C_RED, True)
        ],
    ],
    alignments=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT],
    font_size=13, header_size=13
)

add_callout(
    s2, 0.50, 5.75, 12.33, 1.10,
    "这次汇报只做一件事：回答一个干净的问题",
    "在掌握同样信息（Ni、Zn、位置）的情况下，MOGP（参考身边邻居）预测 Cd 是否比线性回归更准？给出的 95% 误差范围是否更可信？\n"
    "（第一次汇报里的双尺度分解、BIC 选 Q 等内容与本次问题无关，正文中全部删去。）",
    bg_color=BG_RED, accent_color=C_RED, title_size=13.5, body_size=12.5
)

add_footer(s2, f"2 / {TOTAL_MAIN}")


# ==============================================================================
# 第 3 页：问题和数据
# ==============================================================================
s3 = prs.slides.add_slide(BLANK)
set_bg(s3)
add_header(s3, "03", "问题与数据：藏起 100 个点的 Cd 来考模型",
           "瑞士 Jura 土壤数据集（359 个采样点）：测镉（Cd）贵，测镍（Ni）和锌（Zn）便宜")

tb3 = s3.shapes.add_table(4, 3, Inches(0.50), Inches(1.65), Inches(7.55), Inches(2.55)).table
style_table(
    tb3,
    col_widths=[2.15, 2.10, 3.30],
    rows_data=[
        ["变量 / 信息", "哪些点可见", "在考试中的角色"],
        [
            ("Cd（镉，目标变量）", C_RED, True),
            ("259 个可见 / 100 个藏起", C_RED, True),
            "藏起的 100 个点谁都不准偷看，只用来最后打分"
        ],
        [
            "Ni（镍）、Zn（锌）",
            "全部 359 个点可见",
            "便宜的辅助指标，所有模型在测试点上都能直接用"
        ],
        [
            "位置 (X, Y) · 岩石 · 土地用途",
            "全部 359 个点可见",
            "每个采样点的坐标与分类背景（按各模型规则使用）"
        ],
    ],
    alignments=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.LEFT],
    font_size=12.5, header_size=12.5
)

add_callout(
    s3, 0.50, 4.45, 7.55, 1.15,
    "直观比喻：这场考试考的是什么？",
    "测试点（右图橙点）自己的 Ni 和 Zn 是已知的。\n"
    "• 线性回归：只根据“这个点自己的 Ni、Zn 和所处片区”来推算 Cd。\n"
    "• MOGP：还会转头看看“旁边灰点（邻居）的 Cd 实际是多少”。",
    bg_color=BG_BLUE, accent_color=C_BLUE, title_size=13, body_size=12
)

add_callout(
    s3, 0.50, 5.75, 7.55, 1.05,
    "严格沿用第一次汇报的数据处理（防泄漏）",
    "三个金属取 log 后标准化；其中 Cd 的均值和标准差只用 259 个训练点计算，绝不碰 100 个测试点。",
    bg_color=BG_GRAY, accent_color=C_STEEL, title_size=12.5, body_size=12
)

# 右侧放 fig0_split_map.png
add_box(s3, 8.30, 1.60, 4.53, 5.25, fill_color=C_WHITE, line_color=BORDER_G)
add_picture_fit(s3, FIG / "fig0_split_map.png", 8.38, 1.65, 4.37, 4.65)
add_text(s3, 8.40, 6.32, 4.33, 0.45,
         "图 0：10 次随机划分中的第 0 次（灰点 = 训练点，橙点 = 藏起来的测试点）",
         size=10.5, color=C_MUTED, align=PP_ALIGN.CENTER)

add_footer(s3, f"3 / {TOTAL_MAIN}")


# ==============================================================================
# 第 4 页：参赛者（主比较与附加比较）
# ==============================================================================
s4 = prs.slides.add_slide(BLANK)
set_bg(s4)
add_header(s4, "04", "四个参赛模型：主比较与附加比较",
           "让线性回归一步步变强，看“参考邻居（MOGP）”到底在哪一步还能赢")

tb4 = s4.shapes.add_table(5, 4, Inches(0.50), Inches(1.60), Inches(12.33), Inches(3.25)).table
style_table(
    tb4,
    col_widths=[2.45, 4.18, 3.60, 2.10],
    rows_data=[
        ["参赛模型", "公式 / 模型设置", "用到的信息", "比较定位"],
        [
            "回归 A（普通）",
            "Cd = a + b·Ni + c·Zn",
            "只用这个点自己的 Ni、Zn",
            "基础参考"
        ],
        [
            ("回归 B（+ 位置趋势）", C_BLUE, True),
            ("回归 A + d·X + e·Y + f·X² + g·Y² + h·X·Y", C_BLUE, True),
            ("这个点的 Ni、Zn + 整片区域的大范围趋势（倾斜与弯曲）", C_TEXT, True),
            ("主要对手（主比较）", C_BLUE, True)
        ],
        [
            "回归 C（+ 岩石、土地用途）",
            "回归 B + 岩石类型(0/1) + 土地用途(0/1)",
            "回归 B 的全部信息 + 这个点的岩石类型与土地用途",
            ("附加比较（多用信息）", C_RED, True)
        ],
        [
            ("MOGP", C_DARK, True),
            "第一次汇报选定的模型（Q=2），设置完全不变",
            ("这个点的 Ni、Zn + 身边具体邻居的 Cd 和 Ni/Zn", C_TEXT, True),
            ("被测对象", C_DARK, True)
        ],
    ],
    alignments=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER],
    highlight_rows={2: BG_BLUE},
    font_size=12, header_size=12.5
)

# 下方左右两个并排说明框
add_callout(
    s4, 0.50, 5.05, 6.00, 1.80,
    "主比较：回归 B 对 MOGP（同信息公平对决）",
    "• 两者掌握的原始信息最接近（都只知道 Ni、Zn 和坐标 X, Y）。\n"
    "• 为什么回归 B 要加 X²、Y²、X·Y？只用 X、Y 像一块斜木板；加上平方项像一张能弯的纸，能表达“中间高四周低”，为了公平让对手尽量强。\n"
    "• 比喻：回归 B 知道“你住在城市东区”，MOGP 知道“你隔壁邻居家是什么情况”。",
    bg_color=BG_BLUE, accent_color=C_BLUE, title_size=13, body_size=11.5
)

add_callout(
    s4, 6.83, 5.05, 6.00, 1.80,
    "附加比较：回归 C 对 MOGP（多带一本参考书）",
    "• 回归 C 比 MOGP 多知道“岩石类型”和“土地用途”，不算同信息公平对决。\n"
    "• 为什么要比这一组？为了回答一个实际问题：岩石、土地用途这类便宜现成的分类信息，能不能让简单的线性回归追上复杂的 MOGP？\n"
    "• 不管谁赢，都是有实际价值的诚实结论。",
    bg_color=BG_RED, accent_color=C_RED, title_size=13, body_size=11.5
)

add_footer(s4, f"4 / {TOTAL_MAIN}")


# ==============================================================================
# 第 5 页：公平规则 + 赛前定好的判定标准
# ==============================================================================
s5 = prs.slides.add_slide(BLANK)
set_bg(s5)
add_header(s5, "05", "六条公平规则 + 赛前定死的判定标准",
           "跑实验前先把规则和胜负标准写进方案与代码，结果出来后一个字不改")

rules_s5 = [
    ("1", "同一张考卷", "四个模型跑完全相同的 10 次随机划分（seed 0–9），每次藏起同样的 100 个测试点。"),
    ("2", "同样的复习资料", "只看训练集 259 个 Cd；取 log 与标准化只用训练数据计算，所有模型用同一套输入。"),
    ("3", "赛前定死设置", "MOGP 沿用第一次汇报设置不再调参（自检与上次结果相差 < 10⁻¹⁵）；少于 10 个点的小类别赛前规定并入最大类。"),
    ("4", "同样的评分（比两项）", "① 准不准：看 RMSE（平均错多少，越小越好）；② 误差范围可不可信：看 95% 预测区间实际覆盖率（理想值 95%）。"),
    ("5", "配对检验", "每次划分算“回归 RMSE − MOGP RMSE”，对 10 个差值做配对 t 检验，抵消考卷本身的难易波动。"),
    ("6", "先定标准，再看结果", "跑代码前写好下面的判定标准，结果出来后严格照写，绝不挑对 MOGP 有利的说法。"),
]

for idx, (num, r_title, r_desc) in enumerate(rules_s5):
    col_idx = idx % 2
    row_idx = idx // 2
    x0 = 0.50 + col_idx * 6.25
    y0 = 1.60 + row_idx * 1.22
    bg = BG_GRAY if row_idx % 2 == 1 else C_WHITE
    add_box(s5, x0, y0, 6.08, 1.08, fill_color=bg, line_color=BORDER_G)
    add_box(s5, x0 + 0.16, y0 + 0.29, 0.48, 0.48, fill_color=C_STEEL)
    add_text(s5, x0 + 0.16, y0 + 0.29, 0.48, 0.48, num, size=14, color=C_WHITE,
             bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(s5, x0 + 0.78, y0 + 0.10, 5.15, 0.34, r_title, size=13.5, color=C_DARK, bold=True)
    add_text(s5, x0 + 0.78, y0 + 0.44, 5.15, 0.58, r_desc, size=11.5, color=C_TEXT, line_spacing=1.18)

add_callout(
    s5, 0.50, 5.38, 12.33, 1.45,
    "赛前定死的判定标准（写在看结果之前）",
    "• 准不准（主比较 & 附加比较）：MOGP 与回归的 10 个配对差值，必须满足 配对 t 检验 p < 0.05 且平均误差更小，才写“MOGP 比回归更准”；\n"
    "  只要 p ≥ 0.05，一律如实写“两者准确度没有明显差别”。\n"
    "• 误差范围可不可信：哪个模型的 95% 预测区间实际覆盖率更接近 95%，哪个的误差范围就更可信。",
    bg_color=BG_RED, accent_color=C_RED, title_size=14, body_size=12.5
)

add_footer(s5, f"5 / {TOTAL_MAIN}")


# ==============================================================================
# 第 6 页：结果 ①：准不准（表 1 + 图 1 配对折线图）
# ==============================================================================
s6 = prs.slides.add_slide(BLANK)
set_bg(s6)
add_header(s6, "06", "结果 ①：准不准？—— 主比较 MOGP 赢 6.4%，附加比较打平",
           "表 1 为全场核心依据（单位：标准化 log(Cd)）；右图展示 10 张考卷上的配对走势")

tb6 = s6.shapes.add_table(5, 5, Inches(0.50), Inches(1.60), Inches(6.85), Inches(2.65)).table
style_table(
    tb6,
    col_widths=[1.85, 1.55, 1.15, 1.05, 1.25],
    rows_data=[
        ["模型", "平均 RMSE\n(越小越准)", "MOGP 赢\n几次", "配对 t 检验\np 值", "按赛前标准\n判定"],
        [
            "回归 A (Ni, Zn)",
            "0.691 ± 0.047",
            "10 / 10",
            "0.00003",
            "MOGP 更准"
        ],
        [
            ("回归 B (+位置趋势)\n【主要对手】", C_BLUE, True),
            ("0.675 ± 0.050", C_BLUE, True),
            ("10 / 10", C_BLUE, True),
            ("0.0001", C_BLUE, True),
            ("MOGP 更准\n(少错 6.4%)", C_BLUE, True)
        ],
        [
            ("回归 C (+岩石/土地)\n【附加比较】", C_RED, True),
            ("0.637 ± 0.052", C_RED, True),
            ("6 / 10", C_RED, True),
            ("0.33", C_RED, True),
            ("无明显差别\n(仅差 0.9%)", C_RED, True)
        ],
        [
            ("MOGP (+邻居)", C_DARK, True),
            ("0.632 ± 0.054", C_DARK, True),
            "—",
            "—",
            "基准对照"
        ],
    ],
    alignments=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER],
    highlight_rows={2: BG_BLUE, 3: BG_RED},
    font_size=11.5, header_size=11.5
)

add_callout(
    s6, 0.50, 4.42, 6.85, 1.15,
    "1. 主比较（MOGP vs 回归 B）：MOGP 比回归更准",
    "在掌握同样信息（Ni、Zn、位置）下，MOGP 平均少错 6.4%（平均差 0.043，95% CI [0.028, 0.058]），10 次划分全胜，p = 0.0001 < 0.05。",
    bg_color=BG_BLUE, accent_color=C_BLUE, title_size=12.5, body_size=11.5
)

add_callout(
    s6, 0.50, 5.70, 6.85, 1.15,
    "2. 附加比较（MOGP vs 回归 C）：两者准确度没有明显差别",
    "回归 C 加上岩石和土地用途后（0.637）几乎追平 MOGP（0.632），10 次里 MOGP 只赢 6 次，p = 0.33。简单回归加便宜分类信息就能追平复杂模型！",
    bg_color=BG_RED, accent_color=C_RED, title_size=12.5, body_size=11.5
)

# 右侧放 fig1_paired_lines.png
add_box(s6, 7.55, 1.60, 5.28, 5.25, fill_color=C_WHITE, line_color=BORDER_G)
add_picture_fit(s6, FIG / "fig1_paired_lines.png", 7.63, 1.68, 5.12, 4.35)
add_text(s6, 7.68, 6.10, 5.02, 0.65,
         "图 1 读法：每条灰线是一张考卷（同 100 个测试点）。\n"
         "从回归 B 到 MOGP，10 条线全部往下走；但回归 C 到 MOGP 有升有降。",
         size=11, color=C_TEXT, align=PP_ALIGN.LEFT)

add_footer(s6, f"6 / {TOTAL_MAIN}")


# ==============================================================================
# 第 7 页：结果 ②：误差范围可不可信
# ==============================================================================
s7 = prs.slides.add_slide(BLANK)
set_bg(s7)
add_header(s7, "07", "结果 ②：给出的“误差范围”可不可信？（理想值是 95%）",
           "纠正上次汇报的错误说法：线性回归不仅能给预测区间，而且覆盖率比 MOGP 更诚实")

# 顶部大数字对比卡
cards_s7 = [
    ("95.0%", "理想的目标覆盖率", "承诺“95% 把握”，100 个测试点里就该有约 95 个真值落在区间内", BG_GRAY, C_DARK),
    ("94.4% ~ 95.2%", "线性回归 B / C 的实际覆盖率", "几乎正好等于承诺的 95%（其中回归 C 为 95.2%，全场最准）", BG_GREEN, C_GREEN),
    ("92.3%", "MOGP 的实际覆盖率", "区间最窄（2.31），但略微“自信过头”：说有 95% 把握，实际只有约 92%", BG_RED, C_RED),
]
for idx, (num_s, t_s, d_s, bg_c, ac_c) in enumerate(cards_s7):
    x0 = 0.50 + idx * 4.18
    add_box(s7, x0, 1.60, 3.96, 1.55, fill_color=bg_c, line_color=BORDER_G)
    add_box(s7, x0, 1.60, 3.96, 0.07, fill_color=ac_c)
    add_text(s7, x0 + 0.18, 1.72, 3.60, 0.50, num_s, size=24, color=ac_c, bold=True)
    add_text(s7, x0 + 0.18, 2.24, 3.60, 0.32, t_s, size=13, color=C_DARK, bold=True)
    add_text(s7, x0 + 0.18, 2.56, 3.60, 0.50, d_s, size=11, color=C_TEXT, line_spacing=1.18)

tb7 = s7.shapes.add_table(5, 5, Inches(0.50), Inches(3.38), Inches(12.33), Inches(2.25)).table
style_table(
    tb7,
    col_widths=[2.80, 2.60, 1.90, 1.90, 3.13],
    rows_data=[
        ["模型", "95% 预测区间实际覆盖率\n（10 次均值与范围）", "与理想值 95%\n的偏差", "区间平均宽度\n（越窄越精确）", "按赛前标准评价"],
        [
            "回归 A（Ni, Zn）",
            "93.9% （91% – 97%）",
            "−1.1%",
            "2.74",
            "覆盖率接近 95%，区间较宽"
        ],
        [
            "回归 B（+ 位置趋势）主要对手",
            ("94.4% （91% – 98%）", C_GREEN, True),
            ("−0.6%", C_GREEN, True),
            "2.72",
            ("覆盖率非常准（比 MOGP 更可信）", C_GREEN, True)
        ],
        [
            ("回归 C（+ 岩石、土地用途）", C_GREEN, True),
            ("95.2% （90% – 98%）", C_GREEN, True),
            ("+0.2%（最准）", C_GREEN, True),
            "2.53",
            ("全场最接近理想值 95%", C_GREEN, True)
        ],
        [
            ("MOGP（+ 邻居）", C_RED, True),
            ("92.3% （89% – 96%）", C_RED, True),
            ("−2.7%", C_RED, True),
            ("2.31（最窄）", C_BLUE, True),
            ("区间最窄，但覆盖率略低于 95%", C_RED, True)
        ],
    ],
    alignments=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.LEFT],
    highlight_rows={3: BG_GREEN},
    font_size=12, header_size=12
)

add_callout(
    s7, 0.50, 5.82, 12.33, 1.02,
    "上次汇报的反思与更正",
    "上次说“线性回归给不出不确定度”是错的：课本上的最小二乘预测区间公式不仅能给出每个点的 95% 误差范围，而且在 10 次测试中覆盖率（94.4%–95.2%）比 MOGP（92.3%）更老实。",
    bg_color=BG_BLUE, accent_color=C_BLUE, title_size=13, body_size=12
)

add_footer(s7, f"7 / {TOTAL_MAIN}")


# ==============================================================================
# 第 8 页：为什么 MOGP 能赢回归 B？（图 3 邻居错误方向）
# ==============================================================================
s8 = prs.slides.add_slide(BLANK)
set_bg(s8)
add_header(s8, "08", "为什么 MOGP 能赢回归 B？—— 50 米内邻居的残差高度同向",
           "整片几公里的大趋势帮不上忙，真正有用的空间信息藏在“几十米内的邻居”里")

add_callout(
    s8, 0.50, 1.65, 6.20, 1.55,
    "现象 1：整片区域的大范围弯曲趋势（平方项）没用",
    "• 从回归 A 到回归 B，加上位置和平方项后，RMSE 只从 0.691 微降到 0.675。\n"
    "• 只用训练数据做 F 检验：10 次划分里没有一次显著（p 全部 > 0.10，划分 0 的 p = 0.34）。说明几公里尺度的“大趋势”解释不了 Cd 的波动。",
    bg_color=BG_GRAY, accent_color=C_STEEL, title_size=13, body_size=12
)

add_callout(
    s8, 0.50, 3.40, 6.20, 1.85,
    "现象 2：相距 50 米以内的邻居，回归 B 有 73% 的时候错在同一个方向",
    "• 如果回归没有漏掉邻居信息，它在相邻两个点上“同时猜低”或“同时猜高”的比例应该像抛硬币一样，等于 50%。\n"
    "• 实测（右图）：相距 > 50 米的点对确实是 50%；但相距 ≤ 50 米的邻居点对，同向犯错比例高达 73%！",
    bg_color=BG_RED, accent_color=C_RED, title_size=13, body_size=12
)

add_callout(
    s8, 0.50, 5.45, 6.20, 1.38,
    "一句话解释 MOGP 赢回归 B 的机制",
    "回归在一个点上猜低了，它旁边几十米的点多半也猜低了。\n"
    "所以“邻居猜错了多少”本身就是极有用的信息——回归完全不用邻居，而 MOGP 正好把它用上了。",
    bg_color=BG_BLUE, accent_color=C_BLUE, title_size=13, body_size=12
)

# 右侧放 fig3_neighbour_errors.png
add_box(s8, 6.95, 1.65, 5.88, 5.18, fill_color=C_WHITE, line_color=BORDER_G)
add_picture_fit(s8, FIG / "fig3_neighbour_errors.png", 7.05, 1.72, 5.68, 4.45)
add_text(s8, 7.10, 6.22, 5.58, 0.55,
         "图 3：相距 50 米以内的点对，回归 B 有 73% 的时候犯错方向一致，远超抛硬币的 50%（虚线）。",
         size=11, color=C_TEXT)

add_footer(s8, f"8 / {TOTAL_MAIN}")


# ==============================================================================
# 第 9 页：结论 + 局限 + 下一步
# ==============================================================================
s9 = prs.slides.add_slide(BLANK)
set_bg(s9)
add_header(s9, "09", "结论、如实说明的局限与下一步",
           "严格按照第 5 页赛前定死的标准给出结论，不夸大、不隐瞒")

# 左栏：三条结论
add_box(s9, 0.50, 1.60, 6.00, 5.25, fill_color=BG_BLUE, line_color=BORDER_G)
add_box(s9, 0.50, 1.60, 6.00, 0.48, fill_color=C_BLUE)
add_text(s9, 0.68, 1.66, 5.60, 0.38, "一、按赛前标准的正式结论", size=14, color=C_WHITE, bold=True)

add_rich_paragraphs(
    s9, 0.68, 2.22, 5.65, 4.50,
    [
        [("1. 主比较（MOGP 对 回归 B）：MOGP 更准", 13.5, C_BLUE, True)],
        [("在掌握同样信息（Ni、Zn、位置）下，MOGP 平均少错 6.4%（0.632 vs 0.675），10/10 全胜，配对 t 检验 p = 0.0001。\n"
          "赢的原因：抓住了 50 米内近邻的同向残差（73%），而大范围趋势（平方项）无显著帮助。", 12, C_TEXT, False)],
        [("2. 附加比较（MOGP 对 回归 C）：准确度无明显差别", 13.5, C_RED, True)],
        [("简单回归加上“岩石类型 + 土地用途”后（0.637），几乎追平 MOGP（0.632），仅差 0.9%，p = 0.33。\n"
          "启示：便宜现成的分类背景信息，能让简单方法追上复杂模型。", 12, C_TEXT, False)],
        [("3. 误差范围：回归的区间比 MOGP 更可信", 13.5, C_GREEN, True)],
        [("回归 B（94.4%）和回归 C（95.2%）的实际覆盖率非常接近 95%；MOGP 区间最窄，但覆盖率（92.3%）略低于 95%。", 12, C_TEXT, False)],
    ],
    line_spacing=1.22, space_after=8
)

# 右上栏：局限
add_box(s9, 6.80, 1.60, 6.03, 2.85, fill_color=BG_GRAY, line_color=BORDER_G)
add_box(s9, 6.80, 1.60, 6.03, 0.48, fill_color=C_DARK)
add_text(s9, 6.98, 1.66, 5.65, 0.38, "二、汇报里必须照实说明的局限", size=14, color=C_WHITE, bold=True)

add_rich_paragraphs(
    s9, 6.98, 2.20, 5.68, 2.15,
    [
        [("• 随机划分让“参考邻居”占了便宜：", 12.5, C_RED, True),
         ("Jura 数据里不少采样点相距只有几十米（成簇分布）。随机抽 100 个测试点时，测试点身边常常恰好留有训练点，这对 MOGP 有利。", 12, C_TEXT, False)],
        [("• 回归 C 与 MOGP 不是同信息对决：", 12.5, C_DARK, True),
         ("回归 C 多用了岩石和土地用途，MOGP 这次完全没用这两项信息。", 12, C_TEXT, False)],
        [("• MOGP 的预测区间偏窄（92.3% < 95%）：", 12.5, C_DARK, True),
         ("不确定度估计略微自信过头。", 12, C_TEXT, False)],
    ],
    line_spacing=1.22, space_after=6
)

# 右下栏：自然的下一步
add_box(s9, 6.80, 4.65, 6.03, 2.20, fill_color=BG_RED, line_color=BORDER_G)
add_box(s9, 6.80, 4.65, 6.03, 0.48, fill_color=C_RED)
add_text(s9, 6.98, 4.71, 5.65, 0.38, "三、最自然的下一步", size=14, color=C_WHITE, bold=True)

add_rich_paragraphs(
    s9, 6.98, 5.25, 5.68, 1.50,
    [
        [("把“土地用途/岩石”和“参考邻居”结合起来：", 13, C_RED, True)],
        [("现在回归 C 知道土地用途但不知道邻居，MOGP 知道邻居但不知道土地用途。\n"
          "下一步：把岩石和土地用途也给 MOGP（先用分类信息估一个“基础水平”，再用身边邻居修正残差），检验两者结合能不能更准，同时也让 C 与 MOGP 变成完全同信息的公平对决。", 12, C_TEXT, False)],
    ],
    line_spacing=1.22, space_after=4
)

add_footer(s9, f"9 / {TOTAL_MAIN}")


# ==============================================================================
# 第 10 页：致谢 + 备用页导航
# ==============================================================================
s10 = prs.slides.add_slide(BLANK)
set_bg(s10, C_DARK)

add_text(s10, 0.90, 1.75, 11.50, 0.90,
         "谢谢老师与同学！欢迎批评指正",
         size=36, color=C_WHITE, bold=True, font_name=FONT_HEAD)

add_text(s10, 0.90, 2.75, 11.00, 0.55,
         "所有实验代码（01_compare.py, 02_figures.py）与原始表格（results/*.csv）均可一键复现。",
         size=15, color=RGBColor(0xCB, 0xD5, 0xE1))

add_box(s10, 0.90, 3.55, 11.50, 0.03, fill_color=C_STEEL)
add_text(s10, 0.90, 3.78, 11.50, 0.45,
         "已准备好的备用页（Backup Slides，提问时随时翻阅）：",
         size=15, color=RGBColor(0x9D, 0xB2, 0xD9), bold=True)

backups_nav = [
    ("备用页 B1（第 11 页）", "三个回归模型的完整系数表（系数、标准误、t 值、p 值、R²、调整后 R²）"),
    ("备用页 B2（第 12 页）", "回归 B 平方项的训练集 F 检验（10 次划分） + 小样本类别合并规则 + MOGP 设置自检"),
    ("备用页 B3（第 13 页）", "图 2：四个模型的预测值 vs 真实值散点图（划分 0，100 个测试点）"),
    ("备用页 B4（第 14 页）", "图 3b：回归 B 的空间残差地图 + 测试点按“50 米内有无训练邻居”拆分的误差对比"),
]
for idx, (b_tag, b_desc) in enumerate(backups_nav):
    y0 = 4.35 + idx * 0.60
    add_box(s10, 0.90, y0, 2.50, 0.46, fill_color=C_CARD_D, rounded=True)
    add_text(s10, 0.90, y0, 2.50, 0.46, b_tag, size=12, color=RGBColor(0x7F, 0xA8, 0xE0),
             bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(s10, 3.60, y0 + 0.06, 8.80, 0.40, b_desc, size=13, color=C_WHITE)

add_footer(s10, "致谢 / 备用页导航", dark=True)


# ==============================================================================
# 备用页 B1（第 11 页）：三个回归模型的系数表（划分 0）
# ==============================================================================
s11 = prs.slides.add_slide(BLANK)
set_bg(s11)
add_header(s11, "B1", "备用页 1：三个回归模型的拟合系数表（划分 0，n = 259）",
           "数据来源：results/coefficients_seed0.csv —— 土地用途第 1 类（Landuse_1）系数最大且最显著",
           is_backup=True)

# 左表：回归 A 与 回归 B
add_text(s11, 0.50, 1.52, 5.80, 0.35,
         "■ 回归 A（R² = 0.541，调整后 0.538）与 回归 B（R² = 0.571，调整后 0.559）",
         size=12.5, color=C_DARK, bold=True)

tb11_left = s11.shapes.add_table(12, 5, Inches(0.50), Inches(1.88), Inches(5.85), Inches(4.90)).table
style_table(
    tb11_left,
    col_widths=[1.15, 1.25, 1.15, 1.10, 1.20],
    rows_data=[
        ["模型", "变量", "系数 (SE)", "t 值", "p 值"],
        ["回归 A", "截距", "0.007 (0.042)", "0.16", "0.87"],
        ["回归 A", ("Ni", C_BLUE, True), ("0.423 (0.063)", C_BLUE, True), "6.74", ("1.0e-10", C_BLUE, True)],
        ["回归 A", ("Zn", C_BLUE, True), ("0.406 (0.063)", C_BLUE, True), "6.45", ("5.4e-10", C_BLUE, True)],
        ["回归 B", "截距", "−0.067 (0.073)", "−0.91", "0.36"],
        ["回归 B", ("Ni", C_BLUE, True), ("0.400 (0.066)", C_BLUE, True), "6.09", ("4.3e-09", C_BLUE, True)],
        ["回归 B", ("Zn", C_BLUE, True), ("0.434 (0.064)", C_BLUE, True), "6.82", ("6.6e-11", C_BLUE, True)],
        ["回归 B", "X（东-西）", "0.037 (0.054)", "0.68", "0.50"],
        ["回归 B", ("Y（南-北）", C_RED, True), ("0.141 (0.040)", C_RED, True), "3.54", ("0.00047", C_RED, True)],
        ["回归 B", "X²", "0.072 (0.043)", "1.69", "0.092"],
        ["回归 B", "Y²", "−0.006 (0.027)", "−0.22", "0.83"],
        ["回归 B", "X·Y", "0.025 (0.048)", "0.51", "0.61"],
    ],
    alignments=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER],
    highlight_rows={8: BG_RED},
    font_size=10.5, header_size=11
)

# 右表：回归 C
add_text(s11, 6.65, 1.52, 6.18, 0.35,
         "■ 回归 C（R² = 0.652，调整后 R² = 0.635；参照类：Rock_2, Landuse_3）",
         size=12.5, color=C_DARK, bold=True)

tb11_right = s11.shapes.add_table(14, 4, Inches(6.65), Inches(1.88), Inches(6.18), Inches(4.90)).table
style_table(
    tb11_right,
    col_widths=[1.68, 1.65, 1.25, 1.60],
    rows_data=[
        ["回归 C 变量", "系数 (标准误 SE)", "t 值", "p 值（显著性）"],
        ["截距", "−0.388 (0.092)", "−4.20", "3.7e-05 ***"],
        [("Ni", C_BLUE, True), ("0.380 (0.073)", C_BLUE, True), "5.22", ("3.9e-07 ***", C_BLUE, True)],
        [("Zn", C_BLUE, True), ("0.515 (0.062)", C_BLUE, True), "8.35", ("5.1e-15 ***", C_BLUE, True)],
        ["X", "0.031 (0.054)", "0.57", "0.57"],
        [("Y（越往北越高）", C_BLUE, True), ("0.123 (0.037)", C_BLUE, True), "3.34", ("0.00096 ***", C_BLUE, True)],
        ["X²", "0.062 (0.039)", "1.60", "0.11"],
        ["Y²", "−0.015 (0.028)", "−0.53", "0.60"],
        ["X·Y", "−0.022 (0.045)", "−0.49", "0.63"],
        ["Rock_1", "0.297 (0.153)", "1.94", "0.054 ."],
        ["Rock_3", "0.229 (0.118)", "1.95", "0.052 ."],
        ["Rock_5", "0.269 (0.122)", "2.21", "0.028 *"],
        [("Landuse_1（最关键）", C_RED, True), ("0.876 (0.120)", C_RED, True), ("7.32", C_RED, True), ("3.5e-12 ***", C_RED, True)],
        [("Landuse_2", C_RED, True), ("0.289 (0.098)", C_RED, True), "2.94", ("0.0036 **", C_RED, True)],
    ],
    alignments=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER],
    highlight_rows={12: BG_RED, 13: BG_RED},
    font_size=10, header_size=11
)

add_footer(s11, "备用页 B1 / 4")


# ==============================================================================
# 备用页 B2（第 12 页）：回归 B 平方项 F 检验 + 类别合并说明 + MOGP 设置
# ==============================================================================
s12 = prs.slides.add_slide(BLANK)
set_bg(s12)
add_header(s12, "B2", "备用页 2：平方项 F 检验、小类别合并规则与 MOGP 设置",
           "回答老师可能追问的三个技术细节：为什么平方项留着不删？小类别怎么处理？MOGP 改没改参数？",
           is_backup=True)

# 左侧：10 次划分的 F 检验表
add_text(s12, 0.50, 1.55, 5.80, 0.60,
         "1. 回归 B 平方项（X², Y², X·Y）的 F 检验\n"
         "只用训练数据检验（df₁=3, df₂=251）：10 次里 0 次 p < 0.05",
         size=12.5, color=C_DARK, bold=True)

ft_df = pd.read_csv(RES / "ftest_quadratic.csv")
ft_rows = [["划分 (seed)", "F 统计量", "p 值", "是否 < 0.05"]]
for _, r in ft_df.iterrows():
    ft_rows.append([f"划分 {int(r['seed'])}", f"{r['F']:.2f}", f"{r['p']:.2f}", "否"])

tb12 = s12.shapes.add_table(11, 4, Inches(0.50), Inches(2.20), Inches(5.60), Inches(4.60)).table
style_table(
    tb12,
    col_widths=[1.40, 1.40, 1.40, 1.40],
    rows_data=ft_rows,
    alignments=[PP_ALIGN.CENTER] * 4,
    font_size=10.5, header_size=11
)

# 右侧三个说明框
add_callout(
    s12, 6.40, 1.60, 6.43, 1.60,
    "为什么 F 检验不显著，还是保留回归 B 的平方项？",
    "• 这是赛前规则定死的：跑之前承诺回归 B 使用完整二次趋势面（X, Y, X², Y², X·Y）。\n"
    "• 如果看到 F 检验不显著再去删改回归 B 的公式，就违背了“赛前定死参赛方式”的规则。而且保留或删去对 RMSE 影响极小。",
    bg_color=BG_BLUE, accent_color=C_BLUE, title_size=12.5, body_size=11.5
)

add_callout(
    s12, 6.40, 3.35, 6.43, 1.85,
    "2. 岩石与土地用途的编码及小类别合并规则（赛前规定 < 10 点并入最大类）",
    "• 每一类建 0/1 列，并去掉点最多的那一类作为参照基准（避免与常数项完全共线）。\n"
    "• 岩石类型（359 点）：Rock_2 最多（124 点，作参照类）；Rock_4 只有 6 个点（<10），赛前规定并入 Rock_2；保留 Rock_1、Rock_3、Rock_5 三列。\n"
    "• 土地用途（359 点）：Landuse_3 最多（218 点，作参照类）；Landuse_4 只有 8 个点（<10），并入 Landuse_3；保留 Landuse_1、Landuse_2 两列。",
    bg_color=BG_GRAY, accent_color=C_DARK, title_size=12.5, body_size=11.2
)

add_callout(
    s12, 6.40, 5.35, 6.43, 1.48,
    "3. MOGP 的设置与逐位自检",
    "• 直接调用上一级 ../src/models.py 的 TrueLMC (Q=2)，精确 Cholesky 分解，Adam 训练 500 步（lr=0.05，初值 0.08 与 3.0 km）。\n"
    "• 跑出来的 10 次 RMSE 与第一次汇报 stage10_robustness.csv 最大差仅 1.11×10⁻¹⁶。",
    bg_color=BG_RED, accent_color=C_RED, title_size=12.5, body_size=11.5
)

add_footer(s12, "备用页 B2 / 4")


# ==============================================================================
# 备用页 B3（第 13 页）：图 2 预测值 vs 真实值散点图
# ==============================================================================
s13 = prs.slides.add_slide(BLANK)
set_bg(s13)
add_header(s13, "B3", "备用页 3：四个模型的预测值 vs 真实值散点图（划分 0）",
           "横轴为真实的标准化 log(Cd)，纵轴为模型预测值；点越贴近对角虚线，预测越准",
           is_backup=True)

add_box(s13, 0.50, 1.65, 12.33, 4.10, fill_color=C_WHITE, line_color=BORDER_G)
add_picture_fit(s13, FIG / "fig2_pred_vs_true.png", 0.65, 1.75, 12.03, 3.90)

add_callout(
    s13, 0.50, 5.92, 12.33, 0.95,
    "这张图说明了什么？",
    "在划分 0 的 100 个测试点上，回归 A（0.763）和回归 B（0.757）散得较开；加入岩石与土地用途后的回归 C（0.725）向对角线明显收拢，甚至在这一张考卷上比 MOGP（0.731）略微更准一点（10 张考卷里回归 C 赢了 4 次，MOGP 赢了 6 次）。",
    bg_color=BG_BLUE, accent_color=C_BLUE, title_size=12.5, body_size=11.8
)

add_footer(s13, "备用页 B3 / 4")


# ==============================================================================
# 备用页 B4（第 14 页）：图 3b 残差地图 + 为什么换成柱状图
# ==============================================================================
s14 = prs.slides.add_slide(BLANK)
set_bg(s14)
add_header(s14, "B4", "备用页 4：回归 B 的训练残差地图（划分 0，259 个点）",
           "为什么正文第 8 页没有直接放这张地图，而是改成了两根柱子（图 3）？",
           is_backup=True)

add_box(s14, 0.50, 1.60, 5.65, 5.25, fill_color=C_WHITE, line_color=BORDER_G)
add_picture_fit(s14, FIG / "fig3b_residual_map_B.png", 0.60, 1.68, 5.45, 5.05)

add_callout(
    s14, 6.45, 1.60, 6.38, 2.35,
    "为什么把这张残差地图移到备用页？",
    "• 原计划是用这张地图展示“红点（猜低了）扎堆、蓝点（猜高了）扎堆”。\n"
    "• 但画出来以后发现：整个研究区有 5 公里宽，而“挤在一起”的邻居点相距只有几十米（< 0.05 公里），在 5 公里的地图上几乎重叠在一起，肉眼很难看清簇内的同色现象。\n"
    "• 遵照 PPT 选图三原则（不能一眼看清就不放正文），正文改成了直接统计“50 米内点对同向犯错比例（73% vs 50%）”的柱状图。",
    bg_color=BG_BLUE, accent_color=C_BLUE, title_size=13, body_size=12
)

add_callout(
    s14, 6.45, 4.18, 6.38, 2.67,
    "如果老师问：“测试点里到底有多少个点身边 50 米内就有训练点？”",
    "• 划分 0 的 100 个测试点中，有 40 个点在 50 米内就有训练点，60 个点在 50 米外。\n"
    "• 有 50m 近邻的 40 个点：MOGP (0.592) 明显赢过 回归 B (0.640) 和 回归 C (0.634)。\n"
    "• 无 50m 近邻的 60 个点：MOGP (0.811) 与 回归 B (0.826) 接近，且不如知道土地用途的 回归 C (0.780)。\n"
    "• 这直接证实了第 9 页的结论：MOGP 靠近邻加分，回归 C 靠分类背景加分，两者天然互补。",
    bg_color=BG_RED, accent_color=C_RED, title_size=13, body_size=11.5
)

add_footer(s14, "备用页 B4 / 4")

prs.save(OUT_PPTX)
print(f"[OK] Saved {OUT_PPTX} ({len(prs.slides)} slides)")
